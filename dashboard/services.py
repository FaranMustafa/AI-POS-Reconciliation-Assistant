import csv
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Callable, Dict, Iterable, List, Sequence, Tuple

from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from openpyxl import load_workbook

from imports.models import (
    AuditEvent,
    CanonicalTransactionRecord,
    ImportFile,
    Membership,
    ReconciliationIssueRecord,
    ReconciliationRun,
    Tenant,
)
from reconciliation.engine import run_reconciliation
from reconciliation.importers import (
    parse_card_bank_report,
    parse_expense_report,
    parse_fuel_invoice_report,
    parse_lottery_settlement_report,
    parse_pos_daily_report,
    parse_woocommerce_orders,
)
from reconciliation.types import CanonicalTransaction, ReconciliationIssue


Importer = Callable[[str, str], List[CanonicalTransaction]]

IMPORTERS: Dict[str, Importer] = {
    ImportFile.SourceType.POS_DAILY: parse_pos_daily_report,
    ImportFile.SourceType.CARD_BANK: parse_card_bank_report,
    ImportFile.SourceType.WOOCOMMERCE: parse_woocommerce_orders,
    ImportFile.SourceType.FUEL_INVOICE: parse_fuel_invoice_report,
    ImportFile.SourceType.LOTTERY_SETTLEMENT: parse_lottery_settlement_report,
    ImportFile.SourceType.EXPENSE: parse_expense_report,
}


def get_user_tenant(user) -> Tenant:
    membership = Membership.objects.select_related("tenant").filter(user=user).first()
    if membership is None:
        raise ValueError("Your user account is not connected to a tenant.")
    return membership.tenant


def process_uploaded_import(source_type: str, uploaded_file, user=None) -> ReconciliationRun:
    """Single-file entrypoint preserved for backwards-compatibility with existing tests/views."""

    return process_uploaded_run([(source_type, uploaded_file)], user=user)


def process_uploaded_run(
    uploads: Sequence[Tuple[str, object]], user=None
) -> ReconciliationRun:
    """Create a reconciliation run from one or more uploaded source files.

    Each ``uploads`` entry is a ``(source_type, uploaded_file)`` pair. All imports
    are attached to a single :class:`ReconciliationRun` and the engine is given
    the combined transactions so cross-source rules (e.g., missing settlements,
    WooCommerce vs POS) can fire.
    """

    if not uploads:
        raise ValueError("At least one upload is required to start a run.")

    tenant = get_user_tenant(user)
    actor_label = getattr(user, "username", "") or "system"
    import_files: List[ImportFile] = []
    total_transactions = 0

    for source_type, uploaded_file in uploads:
        if source_type not in IMPORTERS:
            raise ValueError(f"Unsupported source type: {source_type}")
        import_file = ImportFile.objects.create(
            tenant=tenant,
            source_type=source_type,
            original_filename=uploaded_file.name,
            content_type=getattr(uploaded_file, "content_type", "")
            or "application/octet-stream",
            file_size=uploaded_file.size,
            status=ImportFile.Status.RECEIVED,
        )
        AuditEvent.objects.create(
            tenant=tenant,
            action=AuditEvent.Action.IMPORT_CREATED,
            import_file=import_file,
            actor_label=actor_label,
            metadata={"source_type": source_type, "filename": uploaded_file.name},
        )
        try:
            content = _read_tabular_upload(uploaded_file)
            transactions = _parse_transactions(source_type, content, import_file)
            _persist_transactions(import_file, transactions)
            import_file.status = ImportFile.Status.PROCESSED
            import_file.error_message = ""
            import_file.save(update_fields=["status", "error_message", "updated_at"])
            AuditEvent.objects.create(
                tenant=tenant,
                action=AuditEvent.Action.IMPORT_PROCESSED,
                import_file=import_file,
                actor_label=actor_label,
                metadata={"transaction_count": len(transactions)},
            )
            total_transactions += len(transactions)
            import_files.append(import_file)
        except Exception as exc:
            import_file.status = ImportFile.Status.FAILED
            import_file.error_message = str(exc)
            import_file.save(update_fields=["status", "error_message", "updated_at"])
            raise

    run = ReconciliationRun.objects.create(
        tenant=tenant,
        status=ReconciliationRun.Status.COMPLETED,
        started_at=timezone.now(),
        completed_at=timezone.now(),
        summary={
            "import_count": len(import_files),
            "transaction_count": total_transactions,
            "source_types": sorted({imp.source_type for imp in import_files}),
        },
    )
    for import_file in import_files:
        run.import_files.add(import_file)
    AuditEvent.objects.create(
        tenant=tenant,
        action=AuditEvent.Action.RUN_CREATED,
        run=run,
        actor_label=actor_label,
        metadata={"import_file_ids": [imp.id for imp in import_files]},
    )

    combined_transactions: List[CanonicalTransaction] = []
    for import_file in import_files:
        combined_transactions.extend(
            _records_to_domain(import_file.transactions.all())
        )

    issues = run_reconciliation(run_id=str(run.id), transactions=combined_transactions)
    _persist_issues(run, issues)
    run.summary = {
        **run.summary,
        "issue_count": len(issues),
    }
    run.save(update_fields=["summary", "updated_at"])
    return run


class PastedTabularUpload:
    def __init__(self, *, name: str, content: str):
        self.name = name
        self.content_type = "text/plain"
        self._content = content.encode("utf-8")
        self.size = len(self._content)

    def read(self):
        return self._content


def _read_tabular_upload(uploaded_file) -> str:
    raw = uploaded_file.read()
    filename = (getattr(uploaded_file, "name", "") or "").lower()
    content_type = getattr(uploaded_file, "content_type", "") or ""
    if filename.endswith(".xlsx") or "spreadsheetml" in content_type:
        return _xlsx_bytes_to_csv(raw)
    return raw.decode("utf-8-sig")


def _xlsx_bytes_to_csv(raw: bytes) -> str:
    workbook = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    sheet = workbook.active
    output = StringIO()
    writer = csv.writer(output)
    for row in sheet.iter_rows(values_only=True):
        values = ["" if value is None else value for value in row]
        if any(str(value).strip() for value in values):
            writer.writerow(values)
    workbook.close()
    return output.getvalue()


def _parse_transactions(
    source_type: str,
    content: str,
    import_file: ImportFile,
) -> List[CanonicalTransaction]:
    importer = IMPORTERS.get(source_type)
    if importer is None:
        raise ValueError(f"Unsupported source type: {source_type}")
    return importer(content, source_import_id=f"import-{import_file.id}")


def _persist_transactions(
    import_file: ImportFile,
    transactions: Iterable[CanonicalTransaction],
) -> None:
    for transaction in transactions:
        CanonicalTransactionRecord.objects.create(
            import_file=import_file,
            source_transaction_id=transaction.id,
            source_type=transaction.source_type,
            business_date=_parse_date(transaction.business_date),
            posted_at=_parse_datetime(transaction.posted_at),
            category=transaction.category,
            tender_type=transaction.tender_type,
            external_ref=transaction.external_ref,
            payee_key=transaction.payee_key,
            description=transaction.description,
            amount=Decimal(transaction.amount),
            metadata=transaction.metadata,
        )


def _persist_issues(run: ReconciliationRun, issues: Iterable[ReconciliationIssue]) -> None:
    for issue in issues:
        ReconciliationIssueRecord.objects.create(
            run=run,
            severity=issue.severity,
            category=issue.category,
            title=issue.title,
            evidence=issue.evidence,
            recommended_action=issue.recommended_action,
            related_transaction_ids=issue.related_transaction_ids,
        )


def _records_to_domain(records: Iterable[CanonicalTransactionRecord]) -> Iterable[CanonicalTransaction]:
    for record in records:
        yield CanonicalTransaction(
            id=record.source_transaction_id,
            source_import_id=f"import-{record.import_file_id}",
            source_type=record.source_type,
            business_date=record.business_date.isoformat(),
            posted_at=record.posted_at.isoformat(),
            category=record.category,
            tender_type=record.tender_type,
            external_ref=record.external_ref,
            payee_key=record.payee_key,
            description=record.description,
            amount=f"{record.amount:.2f}",
            metadata=record.metadata,
        )


def _parse_date(value: str):
    parsed = parse_date(value)
    if parsed is None:
        raise ValueError(f"Invalid date: {value}")
    return parsed


def _parse_datetime(value: str):
    parsed = parse_datetime(value)
    if parsed is None:
        raise ValueError(f"Invalid datetime: {value}")
    if timezone.is_naive(parsed):
        return timezone.make_aware(parsed, timezone=timezone.utc)
    return parsed
